#!/usr/bin/env python3
"""Refresh the supplied labor-market chartbook from BLS into a new DOCX.

Requires Python 3.10+ and python-docx. All downloads use Python's standard library.
No file is evaluated and no code or instructions embedded in a DOCX are executed.
"""
from __future__ import annotations
import argparse
import calendar
import csv
import hashlib
import io
import json
import math
import os
from pathlib import Path
import sys
import time
import urllib.error
import urllib.parse
import urllib.request
from datetime import date, datetime, timezone
from zipfile import ZipFile
from xml.etree import ElementTree as ET

BLS_URL = 'https://api.bls.gov/publicAPI/v2/timeseries/data/'
USER_AGENT = 'LaborChartbookUpdater/1.0 (public economic data)'


def month_index(year, month):
    return int(year) * 12 + int(month) - 1


def month_label(index):
    return f'{index // 12:04d}-{index % 12 + 1:02d}'


def date_index(text):
    value = date.fromisoformat(text + '-01' if len(text) == 7 else text)
    return month_index(value.year, value.month)


def request_bytes(url, body=None, timeout=30):
    headers = {'User-Agent': USER_AGENT}
    if body is not None:
        headers['Content-Type'] = 'application/json'
    req = urllib.request.Request(url, data=body, headers=headers)
    error = None
    for attempt in range(3):
        try:
            with urllib.request.urlopen(req, timeout=timeout) as response:
                return response.read()
        except (urllib.error.URLError, TimeoutError, OSError) as exc:
            error = exc
            if isinstance(exc, urllib.error.HTTPError) and exc.code not in (429, 500, 502, 503, 504):
                break
            if attempt < 2:
                time.sleep(1 + attempt)
    raise RuntimeError(f'Download failed for {url.split("?")[0]}: {error}')


def parse_bls_data(response, requested, first_year, last_year):
    if response.get('status') != 'REQUEST_SUCCEEDED':
        raise RuntimeError('BLS API: ' + '; '.join(response.get('message', ['request failed'])))
    parsed = {sid: {} for sid in requested}
    for series in response.get('Results', {}).get('series', []):
        sid = series.get('seriesID')
        if sid not in parsed:
            continue
        for point in series.get('data', []):
            period = point.get('period', '')
            if not (len(period) == 3 and period.startswith('M') and period[1:].isdigit()):
                continue
            month = int(period[1:])
            year = int(point['year'])
            if not 1 <= month <= 12 or not first_year <= year <= last_year:
                continue  # Exclude annual averages M13 and out-of-window responses.
            try:
                value = float(point['value'].replace(',', ''))
            except (ValueError, TypeError):
                continue
            if not math.isfinite(value):
                continue
            notes = [n for n in point.get('footnotes', []) if n.get('code') or n.get('text')]
            idx = month_index(year, month)
            parsed[sid][idx] = {'value': value, 'footnotes': notes, 'provider': 'BLS'}
    return parsed


def fetch_fred(series, first_year, last_year, raw_dir, reason, log):
    sid = series.get('fred_id')
    if not sid:
        raise RuntimeError(f'No verified FRED mirror for {series["bls_id"]}; retry the BLS download.')
    params = urllib.parse.urlencode({'id': sid, 'cosd': f'{first_year}-01-01', 'coed': f'{last_year}-12-31'})
    raw = request_bytes('https://fred.stlouisfed.org/graph/fredgraph.csv?' + params)
    (raw_dir / f'fred_{sid}_{first_year}_{last_year}.csv').write_bytes(raw)
    reader = csv.DictReader(io.StringIO(raw.decode('utf-8-sig')))
    if sid not in (reader.fieldnames or []):
        raise RuntimeError(f'Invalid FRED CSV for {sid}')
    result = {}
    for row in reader:
        text = row.get('observation_date', row.get('DATE', ''))
        try:
            idx = date_index(text)
            value = float(row[sid])
        except (ValueError, TypeError):
            continue
        if math.isfinite(value) and first_year <= idx // 12 <= last_year:
            result[idx] = {'value': value, 'footnotes': [], 'provider': 'FRED (BLS source)'}
    log.append({'event': 'fred_fallback', 'bls_id': series['bls_id'], 'fred_id': sid, 'reason': reason,
                'note': 'FRED mirror may lag BLS and does not include BLS point footnotes.'})
    print(f'FRED fallback for {series["bls_id"]}: {reason}', file=sys.stderr)
    return result


def fetch_all(config, first_year, last_year, raw_dir, fallback=False):
    series = config['series']
    output = {key: {} for key in series}
    log = []
    keys = list(series)
    # Stay within unregistered BLS limits: 25 series, 10 years per request.
    for first in range(first_year, last_year + 1, 10):
        last = min(first + 9, last_year)
        for offset in range(0, len(keys), 25):
            batch = keys[offset:offset+25]
            ids = [series[k]['bls_id'] for k in batch]
            payload = {'seriesid': ids, 'startyear': str(first), 'endyear': str(last)}
            key = os.environ.get('BLS_API_KEY')
            if key:
                payload['registrationkey'] = key
            error = None
            try:
                raw = request_bytes(BLS_URL, json.dumps(payload).encode())
                response = json.loads(raw)
                (raw_dir / f'bls_{first}_{last}_{offset//25+1}.json').write_text(json.dumps(response, indent=2) + '\n')
                parsed = parse_bls_data(response, ids, first, last)
                log.append({'event': 'bls_request', 'start_year': first, 'end_year': last,
                            'series_count': len(ids), 'messages': response.get('message', [])})
            except (RuntimeError, ValueError) as exc:
                error = str(exc)
                parsed = {}
                if not fallback:
                    raise RuntimeError(error + ' No DOCX was written. Retry later or explicitly enable --allow-fred-fallback.') from exc
            for name in batch:
                sid = series[name]['bls_id']
                values = parsed.get(sid, {})
                if not values:
                    reason = error or f'BLS returned no observations for {sid} during {first}-{last}'
                    if fallback:
                        values = fetch_fred(series[name], first, last, raw_dir, reason, log)
                    else:
                        raise RuntimeError(reason + '. No DOCX was written; source mapping or availability needs review.')
                output[name].update(values)
    return output, log


def derive_table(definition, data):
    names = definition['components']
    dates = set(data[names[0]])
    for name in names[1:]:
        dates.intersection_update(data[name])
    levels = {}
    for idx in dates:
        records = [data[n][idx] for n in names]
        levels[idx] = {
            'value': sum(r['value'] for r in records),
            'footnotes': [note for r in records for note in r.get('footnotes', [])],
            'provider': ', '.join(sorted({r['provider'] for r in records})),
        }
    lag = definition.get('lag_months', 0)
    if not lag:
        return levels
    # Calendar offsets, never previous-row offsets: missing months stay missing.
    return {idx: {'value': point['value'] - levels[idx-lag]['value'],
                  'footnotes': point['footnotes'] + levels[idx-lag]['footnotes'],
                  'provider': ', '.join(sorted({point['provider'], levels[idx-lag]['provider']}))}
            for idx, point in levels.items() if idx-lag in levels}


def format_value(point, digits):
    if point is None:
        return '\u2014'
    value = point['value']
    if abs(value) < 0.5 * 10 ** -digits:
        value = 0.0
    provisional = any(str(n.get('code', '')).upper() == 'P' or 'preliminary' in n.get('text', '').lower()
                      for n in point.get('footnotes', []))
    return f'{value:.{digits}f}' + ('*' if provisional else '')


def inspect_reference(path):
    if path is None:
        return None
    with ZipFile(path) as z:
        root = ET.fromstring(z.read('word/document.xml'))
        images = root.findall('.//{http://schemas.openxmlformats.org/drawingml/2006/main}blip')
    return {'path': str(path), 'sha256': hashlib.sha256(path.read_bytes()).hexdigest(), 'image_count': len(images)}


def make_document(config, results, output, reference, start_year, end_year, retrieved, log):
    from docx import Document
    from docx.shared import Inches, Pt, RGBColor
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    from docx.enum.style import WD_STYLE_TYPE
    from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    doc = Document(str(reference)) if reference else Document()
    # Reuse the reference page system/styles, but replace all old text, pictures,
    # cached TOC and tables, including the title stored in a content control.
    body = doc._element.body
    for child in list(body):
        if child.tag != qn('w:sectPr'):
            body.remove(child)
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.bottom_margin = section.left_margin = section.right_margin = Inches(1)
    for sec in doc.sections:
        for part in (sec.header, sec.footer):
            for p in part.paragraphs:
                p.clear()
    for style_name in ['Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Caption']:
        if style_name not in doc.styles:
            doc.styles.add_style(style_name, WD_STYLE_TYPE.PARAGRAPH)
        style = doc.styles[style_name]
        style.font.name = 'Calibri'
        fonts = style.element.get_or_add_rPr().find(qn('w:rFonts'))
        if fonts is not None:
            for attr in list(fonts.attrib):
                if 'Theme' in attr:
                    del fonts.attrib[attr]
            for script in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
                fonts.set(qn('w:' + script), 'Calibri')
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after = Pt(6)
        rpr = style.element.find(qn('w:rPr'))
        if rpr is not None:
            for color in rpr.findall(qn('w:color')):
                for attr in ['themeColor', 'themeTint', 'themeShade']:
                    color.attrib.pop(qn('w:' + attr), None)
        for border in list(style.element.findall('.//' + qn('w:pBdr'))):
            border.getparent().remove(border)
    doc.styles['Normal'].font.size = Pt(10)
    doc.styles['Normal'].paragraph_format.line_spacing = 1.05
    doc.styles['Title'].font.size = Pt(24)
    doc.styles['Heading 1'].font.size = Pt(15)
    doc.styles['Heading 2'].font.size = Pt(11)
    doc.styles['Caption'].font.size = Pt(8)
    latest = sorted({max(v) for v in results.values() if v})
    latest_text = month_label(latest[-1]) if latest else 'unavailable'
    doc.add_paragraph(config['title'], 'Title')
    doc.add_paragraph(f'Updated {retrieved[:10]}  |  Latest observation {latest_text}', 'Subtitle')
    doc.add_paragraph('Monthly reference tables for employment, the labor force, participation and unemployment. '
                      'Compare published levels and changes by month and year, with the latest observations summarized below.')
    doc.add_heading('How to read the tables', level=1)
    for note in config['notes']:
        doc.add_paragraph(note)
    doc.add_paragraph('An asterisk marks a preliminary observation or a change that uses one. A dash means the observation or comparison month is unavailable; future months remain blank. All values may be revised.')
    if any(e['event'] == 'fred_fallback' for e in log):
        doc.add_paragraph('Some data were downloaded through the FRED mirror of BLS. Source paths are identified below each table and in the run manifest; FRED observations do not carry BLS preliminary footnotes.')
    doc.add_heading('Latest observations', level=1)
    selected = ['nonfarm_level', 'manufacturing_level', 'labor_force_level', 'unemployment', 'lfpr', 'prime_lfpr', 'black', 'hispanic', 'youth']
    for tid in selected:
        definition = next((t for t in config['tables'] if t['id'] == tid), None)
        if not definition or not results[tid]:
            continue
        idx = max(results[tid]); meta = config['series'][definition['components'][0]]
        doc.add_paragraph(f'{definition["title"]}: {format_value(results[tid][idx], meta["digits"])} '
                          f'{meta["units"]} ({month_label(idx)}).')
    doc.add_paragraph('Source: U.S. Bureau of Labor Statistics, Current Employment Statistics and Current Population Survey. '
                      'Series identifiers accompany each table. Downloaded data and complete point footnotes are included with this update.', 'Caption')

    def set_cell(cell, text, header=False, shade=None, year=False):
        cell.text = text
        cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_after = p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.line_spacing = 1
        for run in p.runs:
            run.font.name = 'Calibri'; run.font.size = Pt(7.5)
            run.bold = header or year
            run.font.color.rgb = RGBColor(255, 255, 255) if header else RGBColor(0, 0, 0)
        tcpr = cell._tc.get_or_add_tcPr()
        margins = OxmlElement('w:tcMar')
        for side, size in [('top', 22), ('bottom', 22), ('left', 15), ('right', 15)]:
            node = OxmlElement('w:' + side); node.set(qn('w:w'), str(size)); node.set(qn('w:type'), 'dxa'); margins.append(node)
        tcpr.append(margins)
        fill = OxmlElement('w:shd'); fill.set(qn('w:fill'), '344654' if header else (shade or 'FFFFFF')); tcpr.append(fill)
        borders = OxmlElement('w:tcBorders')
        for side in ['top', 'left', 'bottom', 'right']:
            edge = OxmlElement('w:' + side); edge.set(qn('w:val'), 'single'); edge.set(qn('w:sz'), '4'); edge.set(qn('w:color'), 'D9D9D9'); borders.append(edge)
        tcpr.append(borders)

    panels = []
    for definition in config['tables']:
        for chunk_start in range(start_year, end_year + 1, 18):
            panels.append((definition, chunk_start, min(chunk_start+17, end_year)))
    for panel_number, (definition, first, last) in enumerate(panels):
        if panel_number % 2 == 0:
            doc.add_page_break()
        elif panel_number:
            spacer = doc.add_paragraph(); spacer.paragraph_format.space_after = Pt(7)
        values = results[definition['id']]
        metas = [config['series'][key] for key in definition['components']]
        digits = max(meta['digits'] for meta in metas)
        title = definition['title']
        if first != start_year or last != end_year:
            title += f' {first} to {last}'
        doc.add_heading(title, level=2)
        unit = 'Thousands of jobs' if metas[0]['bls_id'].startswith('CES') else ('Thousands of people' if metas[0]['units'] == 'thousands' else 'Percent')
        adjustment = 'Seasonally adjusted' if metas[0]['seasonal_adjustment'] == 'SA' else 'Not seasonally adjusted'
        latest_idx = max(values) if values else None
        p = doc.add_paragraph(f'{unit}  |  {adjustment}  |  Latest {month_label(latest_idx) if latest_idx is not None else "unavailable"}', 'Caption')
        p.paragraph_format.keep_with_next = True
        table = doc.add_table(rows=1, cols=13)
        table.alignment = WD_TABLE_ALIGNMENT.CENTER; table.autofit = False
        widths = [0.42] + [(6.5-0.42)/12]*12
        for c, width in zip(table.columns, widths): c.width = Inches(width)
        header = ['Year'] + [calendar.month_abbr[m] for m in range(1, 13)]
        for cell, text, width in zip(table.rows[0].cells, header, widths):
            cell.width = Inches(width); set_cell(cell, text, True)
        trpr = table.rows[0]._tr.get_or_add_trPr(); repeat = OxmlElement('w:tblHeader'); trpr.append(repeat)
        for year in range(first, last+1):
            cells = table.add_row().cells
            shade = 'EEF2F5' if (year-first) % 2 == 0 else 'FFFFFF'
            for i, (cell, width) in enumerate(zip(cells, widths)):
                cell.width = Inches(width)
                if i == 0:
                    text = str(year)
                else:
                    idx = month_index(year, i)
                    text = '' if latest_idx is not None and idx > latest_idx else format_value(values.get(idx), digits)
                set_cell(cell, text, shade=shade, year=i == 0)
            no_split = OxmlElement('w:cantSplit'); table.rows[-1]._tr.get_or_add_trPr().append(no_split)
        ids = ' + '.join(m['bls_id'] for m in metas)
        providers = sorted({point['provider'] for point in values.values()})
        p = doc.add_paragraph(f'Source: BLS {ids}. Download: {"; ".join(providers)}.', 'Caption')
        p.paragraph_format.space_before = Pt(3); p.paragraph_format.space_after = Pt(0)
        if definition.get('cps_level_change'):
            p = doc.add_paragraph('January comparisons can include population-control revisions. Changes are differences in published levels.', 'Caption')
            p.paragraph_format.space_after = Pt(0)
    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer.add_run('Economic labor market chartbook  |  ')
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); footer._p.append(field)
    for run in footer.runs: run.font.size = Pt(8)
    doc.core_properties.title = config['title']
    doc.core_properties.subject = 'Updated monthly BLS labor market reference tables'
    doc.core_properties.author = ''; doc.core_properties.last_modified_by = ''
    # Drop orphaned image relationships from the source copy so no old table
    # image survives inside the delivered package after the body is rebuilt.
    image_rels = [rid for rid, rel in doc.part.rels.items() if rel.reltype.endswith('/image')]
    for rid in image_rels: del doc.part.rels[rid]
    doc.save(str(output))


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--input', type=Path, help='Original chartbook DOCX used for style and provenance; never overwritten')
    parser.add_argument('--config', type=Path, default=Path(__file__).with_name('chartbook_config.json'))
    parser.add_argument('--output-dir', type=Path, default=Path('chartbook-output'))
    parser.add_argument('--start-year', type=int, help='First displayed year; default is 2009')
    parser.add_argument('--years', type=int, help='Display only the most recent N calendar years through the cutoff year')
    parser.add_argument('--as-of', help='Observation cutoff YYYY-MM or YYYY-MM-DD; not a historical data vintage')
    parser.add_argument('--allow-fred-fallback', action='store_true', help='Explicitly permit a logged FRED mirror fallback')
    parser.add_argument('--offline-data', type=Path, help='Explicitly rebuild from a previously saved observations.json')
    parser.add_argument('--overwrite', action='store_true', help='Replace prior generated outputs; never the input')
    args = parser.parse_args(argv)
    config = json.loads(args.config.read_text())
    if args.start_year and args.years:
        parser.error('Choose --start-year or --years, not both')
    cutoff = date_index(args.as_of) if args.as_of else month_index(date.today().year, date.today().month)
    last_year = cutoff // 12
    first_year = last_year - args.years + 1 if args.years else args.start_year or config.get('default_start_year', 2009)
    if args.years is not None and args.years < 1 or not 1940 <= first_year <= last_year:
        parser.error('Invalid display window')
    out = args.output_dir.resolve(); out.mkdir(parents=True, exist_ok=True)
    doc_path = out / 'Labor Market Chartbook Updated.docx'
    if args.input and (args.input.resolve() == doc_path.resolve() or args.input.resolve().parent == out):
        parser.error('Use an output directory separate from the input document')
    managed = ['Labor Market Chartbook Updated.docx', 'observations.json', 'observations.csv', 'tables.csv', 'manifest.json']
    if not args.overwrite and any((out / name).exists() for name in managed):
        parser.error('Generated output already exists; choose a new directory or pass --overwrite')
    reference = inspect_reference(args.input)
    if reference and reference['image_count'] != 34:
        parser.error(f'Expected the supplied 34-image chartbook; found {reference["image_count"]} images. Review configuration for a different source document.')
    # Validate config before downloads or authoring.
    if len({t['id'] for t in config['tables']}) != len(config['tables']):
        parser.error('Duplicate table ids in configuration')
    for table in config['tables']:
        if not table['components'] or any(name not in config['series'] for name in table['components']):
            parser.error('Invalid table component mapping')
        metas = [config['series'][name] for name in table['components']]
        if len({(m['units'], m['seasonal_adjustment']) for m in metas}) != 1:
            parser.error('Cannot sum components with different units or seasonal adjustment')
    raw_dir = out / 'raw'; raw_dir.mkdir(exist_ok=True)
    retrieved = datetime.now(timezone.utc).isoformat()
    if args.offline_data:
        snapshot = json.loads(args.offline_data.read_text())
        if snapshot.get('series_metadata') != config['series']:
            parser.error('Offline snapshot series metadata does not match the current configuration')
        retrieved = snapshot['retrieved_at']
        data = {name: {int(k): v for k, v in vals.items()} for name, vals in snapshot['series'].items()}
        log = list(snapshot.get('download_events', [])) + [{'event': 'offline_rebuild', 'path': str(args.offline_data.resolve()), 'original_retrieved_at': retrieved}]
    else:
        print(f'Downloading {len(config["series"])} BLS series for {first_year-1}-{last_year} ...', flush=True)
        data, log = fetch_all(config, first_year-1, last_year, raw_dir, args.allow_fred_fallback)
    data = {name: {idx: point for idx, point in values.items() if idx <= cutoff} for name, values in data.items()}
    if any(not data.get(name) for name in config['series']):
        raise RuntimeError('One or more required series have no observations at or before the cutoff. No DOCX was written.')
    results = {t['id']: {idx: p for idx, p in derive_table(t, data).items() if first_year <= idx//12 <= last_year}
               for t in config['tables']}
    if any(not points for points in results.values()):
        raise RuntimeError('One or more tables have no valid observations. No DOCX was written.')
    manifest = {'generated_at': datetime.now(timezone.utc).isoformat(), 'retrieved_at': retrieved,
                'observation_cutoff': month_label(cutoff), 'display_start_year': first_year,
                'display_end_year': last_year, 'reference': reference,
                'config_sha256': hashlib.sha256(args.config.read_bytes()).hexdigest(),
                'table_count': len(config['tables']), 'series_count': len(config['series']),
                'events': log, 'series': {}, 'tables': {}}
    for name, points in data.items():
        latest = max(points)
        manifest['series'][name] = {**config['series'][name], 'latest_observation': month_label(latest),
                                    'latest_value': points[latest]['value'], 'observations': len(points)}
    for table in config['tables']:
        points = results[table['id']]; latest = max(points)
        gaps = [month_label(idx) for idx in range(month_index(first_year, 1), latest+1) if idx not in points]
        manifest['tables'][table['id']] = {'title': table['title'], 'components': table['components'],
                                          'lag_months': table.get('lag_months', 0),
                                          'latest_observation': month_label(latest),
                                          'latest_value': points[latest]['value'], 'missing_months': gaps}
    snapshot = {'retrieved_at': retrieved, 'series_metadata': config['series'], 'series': data, 'download_events': log}
    (out / 'observations.json').write_text(json.dumps(snapshot, indent=2) + '\n')
    for filename, datasets, is_raw in [('observations.csv', data, True), ('tables.csv', results, False)]:
        with (out / filename).open('w', newline='') as stream:
            writer = csv.writer(stream); writer.writerow(['id', 'date', 'value', 'provider', 'footnotes'])
            for key, points in datasets.items():
                for idx, point in sorted(points.items()):
                    writer.writerow([config['series'][key]['bls_id'] if is_raw else key, month_label(idx),
                                     point['value'], point['provider'], json.dumps(point['footnotes'])])
    temp_doc = out / '.chartbook.pending.docx'
    try:
        make_document(config, results, temp_doc, args.input, first_year, last_year, retrieved, log)
        if args.input and hashlib.sha256(args.input.read_bytes()).hexdigest() != reference['sha256']:
            raise RuntimeError('Input reference changed during the update; generated file was not finalized.')
        temp_doc.replace(doc_path)
    finally:
        if temp_doc.exists(): temp_doc.unlink()
    manifest['output'] = {'file': doc_path.name, 'sha256': hashlib.sha256(doc_path.read_bytes()).hexdigest()}
    (out / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(f'Wrote {doc_path}\n{len(config["tables"])} refreshed tables. Data and manifest: {out}')
    return 0


if __name__ == '__main__':
    try:
        sys.exit(main())
    except (RuntimeError, OSError, ValueError, KeyError) as exc:
        print(f'ERROR: {exc}', file=sys.stderr)
        sys.exit(1)
